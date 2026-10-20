#!/usr/bin/env python3
"""به‌روزرسانی نسخه و فهرست امکانات در راهنمای وردپرس (v1.0.0 → v1.1.1)"""
from docx import Document

PATH = '/home/z/my-project/download/rasadgar-dollar-wordpress-guide.docx'

REPLACEMENTS = [
    ("نسخهٔ ابزار: ۱٫۰٫۰", "نسخهٔ ابزار: ۱٫۱٫۱"),
    ("دستیار معامله‌گر با نمای بازار و عوامل کمّی و واگرایی دلار، پیشنهاد معامله",
     "دستیار معامله‌گر با نمای بازار و واگرایی دلار، دماسنج فشار خرید و فروش از رأی فنی تریدینگ‌ویو، پنجرهٔ تحلیل و پیش‌بینی طلا (LiteFinance و FXStreet) با نقشهٔ مسیر امروز، پیشنهاد معامله"),
]

def patch_paragraph(p):
    full = p.text
    hit = False
    for old, new in REPLACEMENTS:
        if old in full:
            full = full.replace(old, new)
            hit = True
    if hit and p.runs:
        # حفظ قالب‌بندی اولین run و حذف بقیه
        p.runs[0].text = full
        for r in p.runs[1:]:
            r.text = ''
    return hit

doc = Document(PATH)
count = 0

def walk_tables(tables):
    global count
    for t in tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    count += patch_paragraph(p)
                walk_tables(cell.tables)

for p in doc.paragraphs:
    count += patch_paragraph(p)
walk_tables(doc.tables)

doc.save(PATH)
print(f'patched paragraphs: {count}')
